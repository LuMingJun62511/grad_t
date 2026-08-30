"""
从某题型的 output.jsonl 生成审阅 Excel
用法: python gen_review_xlsx.py <题型相对路径>
"""
import json, sys, os, re

if sys.platform == 'win32':
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

BASE = r"C:\Users\LuMin\Desktop\kg_workspace\出题"


def main():
    if len(sys.argv) < 2:
        print("用法: python gen_review_xlsx.py <题型路径>")
        sys.exit(1)

    qtype = sys.argv[1]
    folder = os.path.join(BASE, qtype)
    path = os.path.join(folder, 'output.jsonl')

    rows = []
    with open(path, 'r', encoding='utf-8') as f:
        for line in f:
            line = line.strip()
            if line:
                rows.append(json.loads(line))

    wb = Workbook()
    ws = wb.active
    ws.title = '审阅'
    headers = ['实体', '关系', '标准答案(宾语)', '生成的题目', '生成的答案', '答案一致?', '审阅备注']
    ws.append(headers)
    for col in range(1, 8):
        c = ws.cell(row=1, column=col)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')

    bad = warn = 0
    r_idx = 2
    for r in rows:
        entity, rel, val, raw = r['entity'], r['relation'], r['value'], r['raw']
        q = a = ''
        # 兼容 ```json 代码围栏
        raw_clean = re.sub(r'^```(?:json)?\s*|\s*```$', '', raw.strip(), flags=re.MULTILINE).strip()
        if raw_clean.startswith('{'):
            try:
                obj = json.loads(raw_clean)
                q = obj.get('question', '')
                a = obj.get('answer', '')
            except:
                q = raw_clean
        else:
            q = raw_clean

        # 一致性判定: 判断题(有truth字段)按答案极性; 反向指代(2-1)按答案=实体名; 其他按答案=宾语
        if 'truth' in r:
            affirm_words = ('是', '对', '支持', '有', '搭载', '采用')
            neg_words = ('不', '没', '无')
            is_affirm = any(w in a for w in affirm_words) and not any(w in a for w in neg_words)
            expected_affirm = (r['truth'] == '真')
            consistent = (is_affirm == expected_affirm) if a else False
        elif '反向指代' in qtype:
            consistent = (a.strip() == entity.strip()) if a else False
        elif not val:
            consistent = True  # 整包注入题(2-2~2-5)无自动一致性判定, 靠人工审
        else:
            consistent = (a.strip() == val.strip()) if a else False

        if not a:
            bad += 1
        elif not consistent:
            warn += 1

        ws.cell(row=r_idx, column=1, value=entity)
        ws.cell(row=r_idx, column=2, value=rel)
        ws.cell(row=r_idx, column=3, value=val)
        ws.cell(row=r_idx, column=4, value=q)
        ws.cell(row=r_idx, column=5, value=a)
        ws.cell(row=r_idx, column=6, value='✓' if consistent else '✗' if a else '空')
        ws.cell(row=r_idx, column=7, value='')

        if not a:
            fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
            for col in range(1, 8):
                ws.cell(row=r_idx, column=col).fill = fill
        elif not consistent:
            fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
            for col in range(1, 8):
                ws.cell(row=r_idx, column=col).fill = fill
        r_idx += 1

    widths = [26, 14, 32, 50, 32, 12, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = 'A2'

    out = os.path.join(folder, '审阅.xlsx')
    wb.save(out)
    print(f'{qtype}: {len(rows)} 条 | 答案不一致 {warn} | 答案空 {bad}')
    print(f'审阅文件: {out}')


if __name__ == '__main__':
    main()
